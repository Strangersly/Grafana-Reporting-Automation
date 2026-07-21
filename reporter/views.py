from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import zipfile
from datetime import timedelta
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import LoginView
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.mail import send_mail
from django.core.mail import get_connection as get_email_connection
from django.db import connection as database_connection
from django.db import transaction
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import (
    DashboardForm,
    GrafanaConnectionForm,
    ImportMappingForm,
    ImportUploadForm,
    InvitationAcceptForm,
    InvitationForm,
    ManualRunForm,
    ScheduleForm,
)
from .models import Dashboard, GrafanaConnection, Invitation, ReportArtifact, ReportRun, Schedule, User
from .permissions import admin_required, operator_required
from .services import create_report_run, manual_periods, safe_artifact_path
from .tasks import execute_report_run


class EmailLoginView(LoginView):
    template_name = "registration/login.html"

    def _keys(self):
        email = self.request.POST.get("username", "").strip().lower()
        forwarded = self.request.META.get("HTTP_X_FORWARDED_FOR", "")
        ip = forwarded.split(",", 1)[0].strip() or self.request.META.get("REMOTE_ADDR", "unknown")
        digest = hashlib.sha256(email.encode("utf-8")).hexdigest()
        return [f"login-limit:email:{digest}", f"login-limit:ip:{ip}"]

    def post(self, request, *args, **kwargs):
        limit = settings.LOGIN_RATE_LIMIT_ATTEMPTS
        if any(int(cache.get(key, 0)) >= limit for key in self._keys()):
            return render(
                request,
                self.template_name,
                {"form": self.get_form(), "rate_limited": True},
                status=429,
            )
        return super().post(request, *args, **kwargs)

    def form_invalid(self, form):
        timeout = settings.LOGIN_RATE_LIMIT_SECONDS
        for key in self._keys():
            try:
                cache.incr(key)
            except ValueError:
                cache.set(key, 1, timeout=timeout)
            else:
                cache.touch(key, timeout=timeout)
        return super().form_invalid(form)

    def form_valid(self, form):
        for key in self._keys():
            cache.delete(key)
        return super().form_valid(form)


@login_required
def home(request):
    settings.REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    disk = shutil.disk_usage(settings.REPORT_ROOT)
    context = {
        "active_schedules": Schedule.objects.filter(enabled=True).count(),
        "queued_runs": ReportRun.objects.filter(status=ReportRun.Status.QUEUED).count(),
        "running_runs": ReportRun.objects.filter(status=ReportRun.Status.RUNNING).count(),
        "failed_runs": ReportRun.objects.filter(status__in=[ReportRun.Status.FAILED, ReportRun.Status.PARTIAL]).count(),
        "dashboard_count": Dashboard.objects.filter(enabled=True).count(),
        "recent_runs": ReportRun.objects.select_related("requested_by", "schedule")[:8],
        "recent_artifacts": ReportArtifact.objects.filter(status=ReportArtifact.Status.SUCCEEDED, superseded=False)[:6],
        "disk_used_percent": round((disk.used / disk.total) * 100, 1) if disk.total else 0,
        "disk_free_gb": round(disk.free / (1024 ** 3), 1),
    }
    return render(request, "reporter/home.html", context)


@operator_required
def dashboard_list(request):
    query = request.GET.get("q", "").strip()
    site = request.GET.get("site", "").strip()
    dashboards = Dashboard.objects.all()
    if query:
        dashboards = dashboards.filter(
            Q(name__icontains=query)
            | Q(site__icontains=query)
            | Q(group__icontains=query)
            | Q(category__icontains=query)
        )
    if site:
        dashboards = dashboards.filter(site=site)
    return render(
        request,
        "reporter/dashboard_list.html",
        {
            "dashboards": dashboards,
            "query": query,
            "site": site,
            "sites": Dashboard.objects.order_by("site").values_list("site", flat=True).distinct(),
        },
    )


@operator_required
def dashboard_edit(request, pk=None):
    dashboard = get_object_or_404(Dashboard, pk=pk) if pk else None
    form = DashboardForm(request.POST or None, instance=dashboard)
    if request.method == "POST" and form.is_valid():
        saved = form.save()
        messages.success(request, f"Dashboard {saved.name} saved.")
        return redirect("dashboard_list")
    return render(request, "reporter/form.html", {"form": form, "title": "Edit dashboard" if dashboard else "Add dashboard"})


@operator_required
@require_POST
def dashboard_delete(request, pk):
    dashboard = get_object_or_404(Dashboard, pk=pk)
    name = dashboard.name
    dashboard.delete()
    messages.success(request, f"Dashboard {name} deleted.")
    return redirect("dashboard_list")


@operator_required
@require_POST
def dashboard_bulk(request):
    ids = request.POST.getlist("dashboard_ids")
    action = request.POST.get("action")
    dashboards = Dashboard.objects.filter(pk__in=ids)
    if action == "enable":
        count = dashboards.update(enabled=True)
    elif action == "disable":
        count = dashboards.update(enabled=False)
    else:
        messages.error(request, "Choose a valid bulk action.")
        return redirect("dashboard_list")
    messages.success(request, f"Updated {count} dashboards.")
    return redirect("dashboard_list")


def _parse_import(upload) -> tuple[list[str], list[dict[str, str]]]:
    if upload.name.lower().endswith(".json"):
        parsed = json.loads(upload.read().decode("utf-8-sig"))
        records = parsed.get("dashboards", []) if isinstance(parsed, dict) else parsed
        if not isinstance(records, list) or not all(isinstance(row, dict) for row in records):
            raise ValidationError("JSON must be a list of dashboards or contain a dashboards list.")
        headers = list(dict.fromkeys(key for row in records for key in row.keys()))
        return headers, [{str(key): str(value) if value is not None else "" for key, value in row.items()} for row in records]
    from openpyxl import load_workbook

    workbook = load_workbook(upload, read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        first = next(values)
    except StopIteration as exc:
        raise ValidationError("The workbook is empty.") from exc
    headers = [str(value).strip() if value is not None else "" for value in first]
    if not any(headers):
        raise ValidationError("The first workbook row must contain column headers.")
    rows = []
    for values_row in values:
        row = {headers[index]: str(value).strip() if value is not None else "" for index, value in enumerate(values_row) if index < len(headers) and headers[index]}
        if any(row.values()):
            rows.append(row)
    return [header for header in headers if header], rows


@operator_required
def dashboard_import(request):
    upload_form = ImportUploadForm()
    mapping_form = None
    preview_rows = []
    if request.method == "POST" and request.POST.get("step") == "upload":
        upload_form = ImportUploadForm(request.POST, request.FILES)
        if upload_form.is_valid():
            try:
                headers, rows = _parse_import(upload_form.cleaned_data["file"])
                if not rows:
                    raise ValidationError("No dashboard rows were found.")
                request.session["dashboard_import"] = {"headers": headers, "rows": rows[:1000]}
                mapping_form = ImportMappingForm(headers)
                preview_rows = rows[:5]
            except (ValidationError, ValueError, json.JSONDecodeError) as exc:
                upload_form.add_error("file", exc)
    elif request.method == "POST" and request.POST.get("step") == "apply":
        pending = request.session.get("dashboard_import") or {}
        headers = pending.get("headers", [])
        rows = pending.get("rows", [])
        mapping_form = ImportMappingForm(headers, request.POST)
        preview_rows = rows[:5]
        if mapping_form.is_valid():
            mapping = mapping_form.cleaned_data
            created = updated = skipped = 0
            errors: list[str] = []
            with transaction.atomic():
                for index, row in enumerate(rows, start=2):
                    data = {
                        "name": row.get(mapping["name_column"], "").strip(),
                        "url": row.get(mapping["url_column"], "").strip(),
                        "site": row.get(mapping.get("site_column"), "").strip() or "default",
                        "group": row.get(mapping.get("group_column"), "").strip() or "default",
                        "category": row.get(mapping.get("category_column"), "").strip() or "default",
                    }
                    if not data["name"] or not data["url"]:
                        skipped += 1
                        errors.append(f"Row {index}: name and URL are required.")
                        continue
                    existing = Dashboard.objects.filter(name=data["name"]).first()
                    candidate = existing or Dashboard(name=data["name"])
                    for key, value in data.items():
                        setattr(candidate, key, value)
                    try:
                        candidate.save()
                    except ValidationError as exc:
                        skipped += 1
                        errors.append(f"Row {index}: {exc}")
                        continue
                    created += int(existing is None)
                    updated += int(existing is not None)
            request.session.pop("dashboard_import", None)
            messages.success(request, f"Import complete: {created} created, {updated} updated, {skipped} skipped.")
            if errors:
                messages.warning(request, " ".join(errors[:5]))
            return redirect("dashboard_list")
    return render(
        request,
        "reporter/dashboard_import.html",
        {"upload_form": upload_form, "mapping_form": mapping_form, "preview_rows": preview_rows},
    )


@operator_required
def schedule_list(request):
    return render(request, "reporter/schedule_list.html", {"schedules": Schedule.objects.prefetch_related("dashboards")})


@operator_required
def schedule_edit(request, pk=None):
    schedule = get_object_or_404(Schedule, pk=pk) if pk else None
    form = ScheduleForm(request.POST or None, instance=schedule)
    if request.method == "POST" and form.is_valid():
        saved = form.save(commit=False)
        if not saved.pk:
            saved.created_by = request.user
        saved.save()
        form.save_m2m()
        messages.success(request, f"Schedule {saved.name} saved.")
        return redirect("schedule_list")
    return render(request, "reporter/form.html", {"form": form, "title": "Edit schedule" if schedule else "Add schedule"})


@operator_required
@require_POST
def schedule_delete(request, pk):
    schedule = get_object_or_404(Schedule, pk=pk)
    schedule.delete()
    messages.success(request, "Schedule deleted.")
    return redirect("schedule_list")


@login_required
def run_list(request):
    status = request.GET.get("status", "")
    runs = ReportRun.objects.select_related("requested_by", "schedule")
    if status in ReportRun.Status.values:
        runs = runs.filter(status=status)
    return render(request, "reporter/run_list.html", {"runs": runs[:200], "selected_status": status, "statuses": ReportRun.Status.choices})


@operator_required
def run_create(request):
    form = ManualRunForm(request.POST or None)
    connection = GrafanaConnection.get_solo()
    if request.method == "POST" and form.is_valid():
        if not connection or not connection.credentials_configured:
            form.add_error(None, "An admin must configure Grafana credentials before running reports.")
        else:
            dashboards = Dashboard.objects.filter(enabled=True) if form.cleaned_data["all_dashboards"] else form.cleaned_data["dashboards"]
            try:
                periods = manual_periods(
                    mode=form.cleaned_data["mode"],
                    year=form.cleaned_data.get("year"),
                    month=form.cleaned_data.get("month"),
                    from_date=form.cleaned_data.get("from_date"),
                    to_date=form.cleaned_data.get("to_date"),
                    filename=form.cleaned_data.get("filename", ""),
                    timezone_name=connection.timezone,
                )
                run = create_report_run(
                    periods=periods,
                    dashboards=dashboards,
                    source=ReportRun.Source.MANUAL,
                    preset=form.cleaned_data["mode"],
                    requested_by=request.user,
                    custom_filename=form.cleaned_data.get("filename", ""),
                )
                try:
                    execute_report_run.delay(str(run.id))
                    messages.success(request, "Report run queued.")
                except Exception:
                    run.summary = "The report was created, but the worker queue is unavailable."
                    run.save(update_fields=["summary"])
                    messages.warning(request, run.summary)
                return redirect("run_detail", pk=run.pk)
            except ValueError as exc:
                form.add_error(None, str(exc))
    return render(request, "reporter/run_create.html", {"form": form, "connection_ready": bool(connection and connection.credentials_configured)})


@login_required
def run_detail(request, pk):
    run = get_object_or_404(ReportRun.objects.select_related("requested_by", "schedule"), pk=pk)
    return render(request, "reporter/run_detail.html", {"run": run, "artifacts": run.artifacts.all()})


@login_required
def run_status(request, pk):
    run = get_object_or_404(ReportRun, pk=pk)
    return JsonResponse(
        {
            "id": str(run.id),
            "status": run.status,
            "status_label": run.get_status_display(),
            "progress": run.progress_percent,
            "completed": run.completed_artifacts,
            "failed": run.failed_artifacts,
            "total": run.total_artifacts,
            "summary": run.summary,
            "finished": run.status in {ReportRun.Status.SUCCEEDED, ReportRun.Status.PARTIAL, ReportRun.Status.FAILED, ReportRun.Status.CANCELLED},
        }
    )


@operator_required
@require_POST
def run_cancel(request, pk):
    run = get_object_or_404(ReportRun, pk=pk)
    if run.status == ReportRun.Status.QUEUED:
        run.status = ReportRun.Status.CANCELLED
        run.cancel_requested = True
        run.finished_at = timezone.now()
        run.artifacts.update(status=ReportArtifact.Status.CANCELLED, finished_at=timezone.now())
        run.save(update_fields=["status", "cancel_requested", "finished_at"])
    elif run.status == ReportRun.Status.RUNNING:
        run.cancel_requested = True
        run.save(update_fields=["cancel_requested"])
    messages.success(request, "Cancellation requested.")
    return redirect("run_detail", pk=run.pk)


@login_required
def run_download(request, pk):
    run = get_object_or_404(ReportRun, pk=pk)
    artifacts = run.artifacts.filter(status=ReportArtifact.Status.SUCCEEDED)
    spool = tempfile.SpooledTemporaryFile(max_size=25 * 1024 * 1024)
    with zipfile.ZipFile(spool, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for artifact in artifacts:
            path = safe_artifact_path(artifact.relative_path)
            if path.exists():
                archive.write(path, arcname=artifact.relative_path)
    spool.seek(0)
    return FileResponse(spool, as_attachment=True, filename=f"grafana-report-{run.id}.zip", content_type="application/zip")


@login_required
def report_library(request):
    artifacts = ReportArtifact.objects.filter(status=ReportArtifact.Status.SUCCEEDED, superseded=False).select_related("run")
    year = request.GET.get("year", "")
    month = request.GET.get("month", "")
    site = request.GET.get("site", "")
    if year.isdigit():
        artifacts = artifacts.filter(folder_year=int(year))
    if month.isdigit():
        artifacts = artifacts.filter(folder_month=int(month))
    if site:
        artifacts = artifacts.filter(site=site)
    context = {
        "artifacts": artifacts[:500],
        "years": ReportArtifact.objects.order_by("-folder_year").values_list("folder_year", flat=True).distinct(),
        "sites": ReportArtifact.objects.order_by("site").values_list("site", flat=True).distinct(),
        "selected_year": year,
        "selected_month": month,
        "selected_site": site,
    }
    return render(request, "reporter/report_library.html", context)


@login_required
def artifact_image(request, pk):
    artifact = get_object_or_404(ReportArtifact, pk=pk, status=ReportArtifact.Status.SUCCEEDED)
    path = safe_artifact_path(artifact.relative_path)
    if not path.exists():
        raise Http404("The report image is missing from disk.")
    if settings.USE_X_ACCEL_REDIRECT:
        response = HttpResponse(content_type="image/png")
        response["X-Accel-Redirect"] = f"/protected-reports/{artifact.relative_path}"
        response["Content-Disposition"] = f'inline; filename="{artifact.filename}"'
        return response
    return FileResponse(path.open("rb"), content_type="image/png", filename=artifact.filename)


@admin_required
@require_POST
def artifact_delete(request, pk):
    artifact = get_object_or_404(ReportArtifact, pk=pk)
    path = safe_artifact_path(artifact.relative_path)
    if path.exists():
        path.unlink()
    artifact.delete()
    messages.success(request, "Report image deleted.")
    return redirect("report_library")


@admin_required
def grafana_settings(request):
    instance = GrafanaConnection.get_solo()
    form = GrafanaConnectionForm(request.POST or None, instance=instance)
    if request.method == "POST" and form.is_valid():
        saved = form.save()
        messages.success(request, f"Grafana connection saved for {saved.base_url}.")
        return redirect("grafana_settings")
    return render(request, "reporter/grafana_settings.html", {"form": form, "connection": instance})


@admin_required
def user_list(request):
    return render(request, "reporter/user_list.html", {"users": User.objects.order_by("email")})


@admin_required
def invite_user(request):
    form = InvitationForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        email = form.cleaned_data["email"]
        role = form.cleaned_data["role"]
        Invitation.objects.filter(email=email, accepted_at__isnull=True).update(accepted_at=timezone.now())
        user, _ = User.objects.get_or_create(email=email, defaults={"role": role, "is_active": False})
        user.role = role
        user.is_active = False
        user.invited_at = timezone.now()
        user.set_unusable_password()
        user.save()
        invitation, token = Invitation.issue(email=email, role=role, invited_by=request.user)
        url = f"{settings.PUBLIC_BASE_URL}{reverse('accept_invitation', args=[token])}"
        try:
            sent = send_mail(
                "Your Grafana Reporter invitation",
                f"You were invited to Grafana Reporter. Set your password within 48 hours:\n\n{url}",
                settings.DEFAULT_FROM_EMAIL,
                [email],
                fail_silently=False,
            )
        except Exception:
            sent = 0
        if sent:
            messages.success(request, f"Invitation sent to {email}.")
        else:
            messages.warning(request, "The account was prepared, but SMTP delivery failed. Check email settings and invite the user again.")
        return redirect("user_list")
    return render(request, "reporter/form.html", {"form": form, "title": "Invite user"})


def accept_invitation(request, token):
    try:
        invitation = Invitation.find_valid(token)
    except ValidationError as exc:
        return render(request, "reporter/invitation_invalid.html", {"error": exc.messages[0]}, status=400)
    form = InvitationAcceptForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            invitation = Invitation.objects.select_for_update().get(pk=invitation.pk)
            if invitation.accepted_at or invitation.expires_at <= timezone.now():
                return render(request, "reporter/invitation_invalid.html", {"error": "This invitation is no longer valid."}, status=400)
            user, _ = User.objects.get_or_create(email=invitation.email)
            user.first_name = form.cleaned_data["first_name"]
            user.last_name = form.cleaned_data["last_name"]
            user.role = invitation.role
            user.is_active = True
            user.accepted_at = timezone.now()
            user.set_password(form.cleaned_data["password1"])
            user.save()
            invitation.accepted_at = timezone.now()
            invitation.save(update_fields=["accepted_at"])
        login(request, user)
        messages.success(request, "Your account is ready.")
        return redirect("home")
    return render(request, "reporter/invitation_accept.html", {"form": form, "invitation": invitation})


@admin_required
@require_POST
def user_toggle(request, pk):
    user = get_object_or_404(User, pk=pk)
    if user == request.user:
        messages.error(request, "You cannot deactivate your own account.")
    else:
        user.is_active = not user.is_active
        user.save(update_fields=["is_active"])
        messages.success(request, f"{user.email} is now {'active' if user.is_active else 'inactive'}.")
    return redirect("user_list")


@admin_required
@require_POST
def user_role(request, pk):
    user = get_object_or_404(User, pk=pk)
    role = request.POST.get("role")
    if user == request.user:
        messages.error(request, "You cannot change your own role.")
    elif role not in User.Role.values:
        messages.error(request, "Choose a valid role.")
    else:
        user.role = role
        user.save(update_fields=["role"])
        messages.success(request, f"Role updated for {user.email}.")
    return redirect("user_list")


def health(request):
    checks: dict[str, bool] = {}
    try:
        with database_connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            checks["database"] = cursor.fetchone()[0] == 1
    except Exception:
        checks["database"] = False
    try:
        cache.set("reporter:health", "ok", timeout=10)
        checks["cache"] = cache.get("reporter:health") == "ok"
    except Exception:
        checks["cache"] = False
    try:
        settings.REPORT_ROOT.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=settings.REPORT_ROOT, delete=True):
            checks["report_storage"] = True
    except Exception:
        checks["report_storage"] = False
    heartbeat = cache.get("reporter:worker-heartbeat")
    if heartbeat:
        try:
            checks["worker"] = timezone.now() - datetime_from_iso(heartbeat) < timedelta(minutes=3)
        except ValueError:
            checks["worker"] = False
    else:
        checks["worker"] = False
    ok = all(checks.values())
    return JsonResponse({"status": "ok" if ok else "degraded", "checks": checks}, status=200 if ok else 503)


@admin_required
def external_health(request):
    import urllib.request

    checks = {"grafana": False, "smtp": False}
    grafana = GrafanaConnection.get_solo()
    if grafana:
        try:
            target = urllib.request.Request(grafana.base_url, method="GET")
            with urllib.request.urlopen(target, timeout=8) as response:
                checks["grafana"] = response.status < 500
        except Exception:
            checks["grafana"] = False
    try:
        email_connection = get_email_connection()
        opened = email_connection.open()
        checks["smtp"] = opened is not False
        email_connection.close()
    except Exception:
        checks["smtp"] = False
    ok = all(checks.values())
    return JsonResponse({"status": "ok" if ok else "degraded", "checks": checks}, status=200 if ok else 503)


def datetime_from_iso(value: str):
    from datetime import datetime

    return datetime.fromisoformat(value)
